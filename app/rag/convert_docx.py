# app/rag/convert_docx.py
import os
import logging
import docx

logger = logging.getLogger(__name__)


def read_docx_pages(docx_path: str) -> list[dict]:
    """
    Đọc trực tiếp nội dung file Word (.docx) bằng python-docx.
    Không phụ thuộc vào ứng dụng Microsoft Word hay COM Automation.
    """
    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"Không tìm thấy file Word: {docx_path}")

    filename = os.path.basename(docx_path)
    logger.info(f"[DOCX LOADER] Đang đọc trực tiếp file Word '{filename}'...")

    try:
        doc = docx.Document(docx_path)
        full_text_lines = []

        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                full_text_lines.append(txt)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    full_text_lines.append(row_text)

        full_text = "\n".join(full_text_lines)

        if not full_text.strip():
            logger.warning(f"[DOCX LOADER] File '{filename}' không chứa văn bản chữ.")
            return []

        page_size = 2000
        pages_data = []
        for i in range(0, len(full_text), page_size):
            page_content = full_text[i : i + page_size]
            page_num = (i // page_size) + 1
            pages_data.append(
                {
                    "text": page_content,
                    "metadata": {
                        "source": filename,
                        "page": page_num,
                        "is_scanned": False,
                    },
                }
            )

        logger.info(f"[DOCX LOADER] Đã nạp {len(pages_data)} trang từ '{filename}'.")
        return pages_data

    except Exception as e:
        logger.exception(f"[DOCX LOADER] Lỗi khi đọc file Word '{docx_path}'")
        raise RuntimeError(f"Không thể đọc file Word: {e}") from e
